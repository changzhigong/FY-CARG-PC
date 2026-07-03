import requests
import pandas as pd
import json
import os
from openpyxl import Workbook, load_workbook
import xlrd
from copy import copy
from collections import defaultdict
from datetime import datetime


class DataFetcher:
    """数据获取类"""
    
    def __init__(self, session):
        self.session = session

    def fetch_data(self, warehouse_id: int, date: str) -> bytes:
        """获取仓库数据"""
        try:
            query_url = (
                f'https://xb.fy-carg.com/dmscloud.part/partBill/partBillQueryOemInfoExport/excel?'
                f'generationstartdate={date}&generationenddate={date}'
                f'&warehouseId={warehouse_id}&isLogisticsPartBill=10041002'
            )
            response = self.session.get(query_url, headers={
                'Accept': 'application/json, text/plain, */*',
                'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8,en-GB;q=0.7,en-US;q=0.6',
                'content-type': 'application/json;charset=UTF-8',
                'Refer': 'https://xb.fy-carg.com/',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/134.0.0.0 Safari/537.36 Edg/134.0.0.0',
            })
            response.raise_for_status()
            return response.content
        except requests.RequestException as e:
            raise RuntimeError(f"数据获取失败: {str(e)}")


class ExcelProcessor:
    """Excel文件处理工具类"""
    
    @staticmethod
    def convert_xls_to_xlsx(input_path: str, output_path: str) -> None:
        """转换xls文件为xlsx格式"""
        wb_xls = xlrd.open_workbook(input_path)
        wb_xlsx = Workbook()
        
        # 处理每个工作表
        for sheet_name in wb_xls.sheet_names():
            ws = wb_xlsx.create_sheet(sheet_name)
            sheet = wb_xls.sheet_by_name(sheet_name)
            
            # 批量写入数据
            for row_idx in range(sheet.nrows):
                row_data = [sheet.cell_value(row_idx, col_idx) for col_idx in range(sheet.ncols)]
                ws.append(row_data)
        
        # 删除默认创建的空白工作表
        del wb_xlsx['Sheet']
        wb_xlsx.save(output_path)

    @staticmethod
    def process_workbook(file_path: str, sort_col_index: int = 20) -> None:
        """处理Excel文件内容"""
        wb = load_workbook(file_path)
        ws = wb.active
        
        if ws is not None:
            # 获取所有数据行（跳过标题行）
            headers = list(ws.values)[0] if ws.values is not None else []
            data_rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                data_rows.append(row)

            # 获取拣货单号列的索引
            sort_col_index = 20  # 根据实际列索引调整
            
            # 改进的排序逻辑（处理空值）
            def sort_key(row):
                value = row[sort_col_index]
                if value is None:  # 空值处理
                    return ('',)  # 空字符串作为排序键
                try:
                    return (str(value),)  # 统一转换为字符串比较
                except:
                    return ('',)
            
            # 按拣货单号排序（空值排在最前）
            sorted_rows = sorted(data_rows, key=sort_key)
            
            # 清空并重新写入数据
            ws.delete_rows(1, ws.max_row)
            ws.append(headers)
            for row in sorted_rows:
                ws.append(row)
            
            # 格式调整
            ExcelProcessor._adjust_format(ws)
            
            # 计算并添加合计
            ExcelProcessor._add_summary(ws, column='M')
        
        wb.save(file_path)

    @staticmethod
    def _adjust_format(ws) -> None:
        """调整工作表格式"""
        # 列宽设置
        column_widths = {'B': 22,'H':10, 'M': 10, 'U': 22}
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        # 隐藏列
        for col in ['E', 'F', 'G','I','J','K','L','V','W','X']:
            ws.column_dimensions[col].hidden = True
        
        # 添加边框
        from openpyxl.styles import Border, Side
        thin_border = Border(
            left=Side(style='thin'), 
            right=Side(style='thin'),
            top=Side(style='thin'), 
            bottom=Side(style='thin')
        )
        for row in ws.iter_rows():
            for cell in row:
                cell.border = thin_border

    @staticmethod
    def _add_summary(ws, column: str) -> None:
        """添加列合计"""
        values = [cell.value for cell in ws[column][1:] if isinstance(cell.value, (int, float))]
        sum_row = len(values) + 1
        ws[f'{column}{sum_row+1}'] = sum(values)